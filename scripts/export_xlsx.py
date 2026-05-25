"""Export the merged ``targeting_profiles_final.json`` to a styled XLSX.

The XLSX is the primary deliverable for the customer (defense commercial).
Two sheets:
  1. ``Eurosatory 2026`` — all 2580 rows, columns ordered for 10-second
     readability, completeness score colour-coded.
  2. ``Long tail`` — the worst 60 (non-manual) rows, for the next manual
     review batch.

Generates BOTH the French canonical file (``Eurosatory_2026_targeting.xlsx``)
and an English-localized variant (``Eurosatory_2026_targeting_en.xlsx``)
using the ``*_en`` fields already present in the JSON.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

EXPORT_DIR = ROOT / "data" / "exports"

# Visible columns. The "(filtres)" columns are the canonical-taxonomy
# columns the user filters on in Excel ; the un-suffixed columns hold
# the specific labels (good for reading).
COLUMNS_FR = [
    ("Société",                   30),
    ("Pays",                      12),
    ("Pavillon",                  10),
    ("Site",                      28),
    ("Activité (1 ligne)",        56),
    ("Produits",                  50),
    ("Catégories produits (filtres)", 38),
    ("Services",                  36),
    ("Catégories services (filtres)", 30),
    ("Cibles clients",            26),
    ("Technologies",              30),
    ("Catégories technos (filtres)", 28),
    ("Pourquoi cibler",           60),
    ("Score",                      7),
    ("Source",                    10),
]

COLUMNS_EN = [
    ("Company",                   30),
    ("Country",                   12),
    ("Pavilion",                  10),
    ("Website",                   28),
    ("Activity (1 line)",         56),
    ("Products",                  50),
    ("Product categories (filters)", 38),
    ("Services",                  36),
    ("Service categories (filters)", 30),
    ("Target customers",          26),
    ("Technologies",              30),
    ("Technology categories (filters)", 28),
    ("Why target",                60),
    ("Score",                      7),
    ("Source",                    10),
]

COLUMNS = COLUMNS_FR  # back-compat default for any caller that imports COLUMNS


def _write_sheet(ws, records, columns=COLUMNS_FR, lang: str = "fr"):
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for i, (name, w) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill(start_color="0F1B2D", end_color="0F1B2D",
                             fill_type="solid")
        c.alignment = Alignment(horizontal="left", vertical="center",
                                wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    # Column indices for special formatting (1-based):
    # 1=Société 5=Activité 6=Produits 7=ProdCat 8=Services 9=SvcCat
    # 10=Cibles 11=Technos 12=TechCat 13=Pourquoi 14=Score 15=Source
    SCORE_COL = 14
    SOURCE_COL = 15
    CATEGORY_COLS = {7, 9, 12}

    # Pick the language-specific list keys when EN, fall back to FR.
    def _g(r, fr_key, en_key):
        if lang == "en":
            v = r.get(en_key)
            if v:
                return v
        return r.get(fr_key)

    for row_idx, r in enumerate(records, start=2):
        # Text values (FR vs EN swap)
        activity = _g(r, "activity_1liner", "activity_1liner_en") or ""
        products = _g(r, "products", "products_en") or []
        services = _g(r, "services", "services_en") or []
        target_buyers = _g(r, "target_buyers", "target_buyers_en") or []
        technologies = _g(r, "technologies", "technologies_en") or []
        why_target = _g(r, "why_target", "why_target_en") or ""

        vals = [
            r["company_name"],
            r["country"] or "",
            r["pavilion"] or "",
            r["website"] or "",
            activity,
            " · ".join(products) if products else "—",
            " · ".join(r.get("products_categories") or []) or "—",
            " · ".join(services) if services else "—",
            " · ".join(r.get("services_categories") or []) or "—",
            ", ".join(target_buyers) if target_buyers else "—",
            " · ".join(technologies) if technologies else "—",
            " · ".join(r.get("technologies_categories") or []) or "—",
            why_target,
            r["completeness_score"],
            r["data_source_strength"],
        ]
        for col_idx, v in enumerate(vals, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = border
            if col_idx == SCORE_COL:
                s = r["completeness_score"]
                if s >= 90:
                    c.fill = PatternFill(start_color="C8F7C5",
                                         end_color="C8F7C5",
                                         fill_type="solid")
                elif s >= 60:
                    c.fill = PatternFill(start_color="FFF4C2",
                                         end_color="FFF4C2",
                                         fill_type="solid")
                else:
                    c.fill = PatternFill(start_color="FFD6CC",
                                         end_color="FFD6CC",
                                         fill_type="solid")
                c.alignment = Alignment(horizontal="center", vertical="center",
                                        wrap_text=True)
                c.font = Font(bold=True)
            elif col_idx == SOURCE_COL and r["data_source_strength"] == "manual":
                c.fill = PatternFill(start_color="DDE9FF",
                                     end_color="DDE9FF",
                                     fill_type="solid")
                c.font = Font(italic=True, color="0F1B2D")
            elif col_idx in CATEGORY_COLS:
                c.fill = PatternFill(start_color="F4F8FB",
                                     end_color="F4F8FB",
                                     fill_type="solid")
                c.font = Font(size=10, color="2A3A4D")
        ws.row_dimensions[row_idx].height = 90


def _enable_autofilter(ws, n_rows: int, n_cols: int) -> None:
    """Enable the Excel AutoFilter (filter dropdowns) on the data area."""
    last_col = openpyxl.utils.get_column_letter(n_cols)
    ws.auto_filter.ref = f"A1:{last_col}{n_rows + 1}"


def _build_workbook(final, long_tail, *, lang: str, main_title: str,
                    lt_title: str):
    cols = COLUMNS_EN if lang == "en" else COLUMNS_FR
    wb = openpyxl.Workbook()
    ws_main = wb.active
    ws_main.title = main_title
    _write_sheet(ws_main, final, columns=cols, lang=lang)
    _enable_autofilter(ws_main, len(final), len(cols))

    ws_lt = wb.create_sheet(lt_title)
    _write_sheet(ws_lt, long_tail, columns=cols, lang=lang)
    _enable_autofilter(ws_lt, len(long_tail), len(cols))
    return wb


def _write_csv(records, *, path: Path, lang: str) -> None:
    """Flat CSV companion (used by the UI's "CSV livrable" button).
    Mirrors the XLSX column order for consistency between the two
    formats — FR or EN depending on ``lang``.
    """
    import csv as _csv
    cols = COLUMNS_EN if lang == "en" else COLUMNS_FR
    headers = [name for name, _ in cols]

    def _g(r, fr_key, en_key):
        if lang == "en":
            v = r.get(en_key)
            if v:
                return v
        return r.get(fr_key)

    with path.open("w", encoding="utf-8", newline="") as f:
        w = _csv.writer(f)
        w.writerow(headers)
        for r in records:
            products = _g(r, "products", "products_en") or []
            services = _g(r, "services", "services_en") or []
            target_buyers = _g(r, "target_buyers", "target_buyers_en") or []
            technologies = _g(r, "technologies", "technologies_en") or []
            w.writerow([
                r["company_name"],
                r["country"] or "",
                r["pavilion"] or "",
                r["website"] or "",
                _g(r, "activity_1liner", "activity_1liner_en") or "",
                " · ".join(products),
                " · ".join(r.get("products_categories") or []),
                " · ".join(services),
                " · ".join(r.get("services_categories") or []),
                ", ".join(target_buyers),
                " · ".join(technologies),
                " · ".join(r.get("technologies_categories") or []),
                _g(r, "why_target", "why_target_en") or "",
                r.get("completeness_score") or 0,
                r.get("data_source_strength") or "",
            ])


def main() -> int:
    final = json.loads(
        (EXPORT_DIR / "targeting_profiles_final.json").read_text(encoding="utf-8")
    )
    long_tail = json.loads(
        (EXPORT_DIR / "profiles_long_tail.json").read_text(encoding="utf-8")
    )

    # Sort the main sheet by score desc then country
    final.sort(key=lambda r: (-r["completeness_score"], r["country"] or "z"))

    # ── French canonical deliverable ──────────────────────────────────
    wb_fr = _build_workbook(
        final, long_tail, lang="fr",
        main_title="Eurosatory 2026",
        lt_title="Long tail (à retraiter)",
    )
    out_fr = EXPORT_DIR / "Eurosatory_2026_targeting.xlsx"
    wb_fr.save(out_fr)
    print(f"Wrote {out_fr}  ({out_fr.stat().st_size:,} bytes)")
    print(f"  - Sheet 1: {len(final)} sociétés (triées par score, autofilter ON)")
    print(f"  - Sheet 2: {len(long_tail)} fiches à retraiter (long tail)")

    # CSV companion (FR) — flat columns matching the XLSX layout.
    csv_fr = EXPORT_DIR / "targeting_profiles_final.csv"
    _write_csv(final, path=csv_fr, lang="fr")
    print(f"Wrote {csv_fr}  ({csv_fr.stat().st_size:,} bytes)")

    # ── English-localized deliverable ─────────────────────────────────
    wb_en = _build_workbook(
        final, long_tail, lang="en",
        main_title="Eurosatory 2026",
        lt_title="Long tail (to retreat)",
    )
    out_en = EXPORT_DIR / "Eurosatory_2026_targeting_en.xlsx"
    wb_en.save(out_en)
    print(f"Wrote {out_en}  ({out_en.stat().st_size:,} bytes)")
    print(f"  - Sheet 1: {len(final)} companies (sorted by score, autofilter ON)")
    print(f"  - Sheet 2: {len(long_tail)} long-tail rows to retreat")

    # CSV companion (EN).
    csv_en = EXPORT_DIR / "targeting_profiles_final_en.csv"
    _write_csv(final, path=csv_en, lang="en")
    print(f"Wrote {csv_en}  ({csv_en.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
